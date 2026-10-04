# Κατάλογος (Επαρχίες & Στοές), ρόλοι ΕπΜΔ/ΕπΜΓρ. και επιλογή παραληπτών από τον Κατάλογο.
import re
from datetime import date, timedelta
from urllib.parse import unquote


def _set_roles(admin, app_module, short, **roles):
    p = [x for x in app_module.provinces_all() if x['short'] == short][0]
    data = {k: str(p.get(k) if p.get(k) is not None else '') for k in app_module.PROVINCE_COLS}
    data.update(roles)
    assert admin.post(f"/provinces/edit/{p['id']}", data=data).status_code == 303
    return p['id']


def test_province_roles_and_form(admin, app_module):
    pid = _set_roles(admin, app_module, 'ΠΜΣτ. Κύπρου', master_name='Ανδρέας Κυπριανός', master_email='gm.cy@example.com',
                     secretary_name='Νίκος Γραμματικός', secretary_email='')
    p = app_module.province_get(pid)
    gm, gs = app_module.province_roles(p)
    assert gm['abbr'] == 'ΠερΜΔ' and gm['email'] == 'gm.cy@example.com' and 'Ανδρέας Κυπριανός' in gm['addressee']
    assert gs['abbr'] == 'ΠερΜΓρ.' and gs['email'] == 'dglcyprus@nglgreece.gr'  # κενό → email Γραμματείας
    t = admin.get(f'/provinces/edit/{pid}').text
    assert 'name="secretary_email"' in t and 'role-pick' in t  # επιλογή από το Μητρώο Μελών
    t = admin.get('/provinces').text
    assert 'Ανδρέας Κυπριανός' in t and 'gm.cy@example.com' in t
    assert admin.post(f'/provinces/edit/{pid}', data=dict({k: str(p.get(k) or '') for k in app_module.PROVINCE_COLS},
                                                          secretary_email='λάθος')).status_code == 400


def test_directory_page(admin, app_module):
    _set_roles(admin, app_module, 'ΕπΜΣτ. Αθηνών', master_name='Γεώργιος Αθηναίος', master_email='gm.ath@example.com',
               secretary_name='Ιωάννης Γραφεύς', secretary_email='gs.ath@example.com')
    t = admin.get('/directory').text
    for s in ['📇 Κατάλογος', 'Γεώργιος Αθηναίος', 'gm.ath@example.com', 'gs.ath@example.com', 'athens.secretary@nglgreece.gr',
              'LA PAIX', 'Όλοι οι ΕπΜΔ', 'data-copy="@visible"']:
        assert s in t, s
    href = re.search(r'href="(/new\?to_name=[^"]*gm\.ath[^"]*)"', t.replace('%40', '@')).group(1)
    letter = admin.get(href.replace('&amp;', '&')).text
    assert 'gm.ath@example.com' in letter and 'Γεώργιος Αθηναίος' in letter  # «Επιστολή προς…» με συμπληρωμένο παραλήπτη
    assert 'href="/directory"' in admin.get('/').text


def test_provinces_export(admin):
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(admin.get('/provinces/export.xlsx').content)).active
    head = [c.value for c in ws[1]]
    assert head[5:9] == ['ΕπΜΔ', 'Email ΕπΜΔ', 'ΕπΜΓρ.', 'Email ΕπΜΓρ.']
    assert any(r[0].value == 'ΕπΜΣτ. Αθηνών' and r[6].value == 'gm.ath@example.com' for r in ws.iter_rows(min_row=2))


def test_letter_recipient_options_include_roles(admin):
    t = admin.get('/new').text
    assert 'gm.ath@example.com' in t and 'Όλοι οι Επαρχιακοί Μεγάλοι Διδάσκαλοι' in t
    assert 'ΕπΜΔ Επαρχιακής Μεγάλης Στοάς Αθηνών, Γεώργιος Αθηναίος' in t


def test_picker_in_visit_emails(admin):
    admin.post('/lodges/new', data={'number': '985', 'name': 'ΚΑΤΑΛΟΓΙΚΗ', 'email': 'lodge985@example.com', 'provincial': 'ΕπΜΣτ. Αθηνών', 'status': 'Ενεργή'})
    admin.post('/reps/new', data={'name': 'Μάρκος', 'surname': 'Καταλογίδης', 'office': 'Μέγας Ευχέτης', 'email': 'markos@example.com'})
    rid = re.search(r'/reps/edit/(\d+)', admin.get('/reps', params={'q': 'Καταλογίδης'}).text).group(1)
    d = (date.today() + timedelta(days=9)).isoformat()
    admin.post('/visits/new', data={'visit_date': d, 'lodge': 'ΚΑΤΑΛΟΓΟΥ', 'lodge_number': '', 'location': '', 'province': 'ΕπΜΣτ. Αθηνών', 'rep_id': rid, 'notes': ''})
    vid = re.search(r'/visits/edit/(\d+)', admin.get('/visits', params={'q': 'ΚΑΤΑΛΟΓΟΥ'}).text).group(1)
    t = admin.get(f'/visits/brief?ids={vid}').text
    assert 'Προσθήκη παραλήπτη από τον Κατάλογο' in t and 'gm.ath@example.com' in t
    assert 'Στοές — ΕπΜΣτ. Αθηνών' in t and 'lodge985@example.com' in t


def test_directory_access(admin):
    from fastapi.testclient import TestClient
    from conftest import APP, login
    admin.post('/users', data={'email': 'plain@example.com', 'role': 'editor', 'active': '1'})
    c = login(TestClient(APP.app, follow_redirects=False), email='plain@example.com')
    assert c.get('/directory').status_code in (200, 403)
    assert 'Edit</a>' not in (c.get('/directory').text if c.get('/directory').status_code == 200 else '')
