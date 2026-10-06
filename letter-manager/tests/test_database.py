# Σελίδα «Βάση Δεδομένων»: όλοι οι πίνακες, αναζήτηση, επεξεργασία πεδίων, Excel και Βιβλίο Πρωτοκόλλου.
from io import BytesIO

from openpyxl import load_workbook


def test_database_home_lists_all_tables(admin, app_module):
    t = admin.get('/database').text
    for k, s in app_module.DB_TABLES.items():
        assert f'href="/database/{k}"' in t and s['label'] in t
    assert 'Βιβλίο Πρωτοκόλλου' in t


def test_every_table_page_and_export_open(admin, app_module):
    for k in app_module.DB_TABLES:
        assert admin.get(f'/database/{k}').status_code == 200, k
        assert admin.get(f'/database/{k}?q=α').status_code == 200, k
        r = admin.get(f'/database/{k}/export.xlsx')
        assert r.status_code == 200 and load_workbook(BytesIO(r.content)).active.max_row >= 1, k
    assert admin.get('/database/nosuch').status_code == 404


def test_search_and_generic_edit(admin, app_module):
    t = admin.get('/database/lodges?q=paix').text
    assert 'LA PAIX' in t
    with app_module.con() as c:
        lid = c.execute("SELECT id FROM lodges ORDER BY id LIMIT 1").fetchone()['id']
        before = dict(c.execute('SELECT * FROM lodges WHERE id=?', (lid,)).fetchone())
    form = admin.get(f'/database/lodges/{lid}').text
    assert 'name="ritual"' in form and 'name="id"' not in form
    data = {k: '' if v is None else str(v) for k, v in before.items() if k not in app_module.DB_READONLY_COLS}
    data['meeting_place'] = 'Δοκιμαστικός Ναός'
    assert admin.post(f'/database/lodges/{lid}', data=data).status_code == 303
    with app_module.con() as c:
        after = dict(c.execute('SELECT * FROM lodges WHERE id=?', (lid,)).fetchone())
    assert after['meeting_place'] == 'Δοκιμαστικός Ναός' and after['name'] == before['name'] and after['number'] == before['number']
    data['number'] = 'όχι αριθμός' if isinstance(before['number'], int) else data['number']
    if isinstance(before['number'], int):
        assert admin.post(f'/database/lodges/{lid}', data=data).status_code == 400
    data['meeting_place'] = before['meeting_place'] or ''
    data['number'] = str(before['number'])
    admin.post(f'/database/lodges/{lid}', data=data)


def test_readonly_tables_cannot_be_edited(admin, app_module):
    with app_module.con() as c:
        r = c.execute('SELECT id FROM letters LIMIT 1').fetchone()
    assert admin.get(f"/database/letters/{r['id'] if r else 1}").status_code == 404
    assert admin.post(f"/database/letters/{r['id'] if r else 1}", data={'subject': 'x'}).status_code == 404


def test_protocol_book_and_system_check(admin):
    t = admin.get('/database/protocol').text
    assert 'Βιβλίο Πρωτοκόλλου' in t
    r = admin.get('/database/protocol?export=1')
    assert r.status_code == 200 and load_workbook(BytesIO(r.content)).active['A1'].value == 'Αρ. Πρωτοκόλλου'
    t = admin.get('/system/check').text
    assert 'Έλεγχος σελίδων' in t and 'Βάση δεδομένων' in t


def test_database_is_admin_only(anon):
    assert anon.get('/database').status_code in (302, 303, 307)
    assert anon.get('/system/check').status_code in (302, 303, 307)


def test_backup_and_restore_roundtrip(admin, app_module):
    import gzip, json
    r = admin.get('/database/backup.json.gz')
    assert r.status_code == 200
    data = json.loads(gzip.decompress(r.content))
    assert data['format'] == 'nglg-backup/1' and 'lodges' in data['tables'] and 'otps' not in data['tables']
    with app_module.con() as c:
        n_lodges = c.execute('SELECT COUNT(*) n FROM lodges').fetchone()['n']
        lid = c.execute('SELECT id FROM lodges ORDER BY id LIMIT 1').fetchone()['id']
        c.execute("UPDATE lodges SET meeting_place='ΑΛΛΑΓΗ ΜΕΤΑ ΤΟ ΑΝΤΙΓΡΑΦΟ' WHERE id=?", (lid,))
        c.execute("DELETE FROM lodges WHERE id=(SELECT MAX(id) FROM lodges)")
    assert admin.get('/database/restore').status_code == 200
    bad = admin.post('/database/restore', data={'confirm': 'όχι'}, files={'file': ('b.json.gz', r.content, 'application/gzip')})
    assert bad.status_code == 400
    bad = admin.post('/database/restore', data={'confirm': 'ΕΠΑΝΑΦΟΡΑ'}, files={'file': ('x.json', b'{}', 'application/json')})
    assert bad.status_code == 400
    ok = admin.post('/database/restore', data={'confirm': 'ΕΠΑΝΑΦΟΡΑ'}, files={'file': ('b.json.gz', r.content, 'application/gzip')})
    assert ok.status_code == 303
    with app_module.con() as c:
        assert c.execute('SELECT COUNT(*) n FROM lodges').fetchone()['n'] == n_lodges
        assert c.execute('SELECT meeting_place FROM lodges WHERE id=?', (lid,)).fetchone()['meeting_place'] != 'ΑΛΛΑΓΗ ΜΕΤΑ ΤΟ ΑΝΤΙΓΡΑΦΟ'
    # μετά την επαναφορά οι νέες εγγραφές παίρνουν νέο αριθμό (σωστή συνέχεια αρίθμησης και στο Postgres)
    assert admin.get('/directory').status_code == 200
    with app_module.con() as c:
        top = c.execute('SELECT MAX(id) m FROM lodges').fetchone()['m']
        new = c.execute("INSERT INTO lodges(number,name) VALUES('999','ΔΟΚΙΜΗ ΜΕΤΑ ΤΗΝ ΕΠΑΝΑΦΟΡΑ')").lastrowid
        assert new > top
        c.execute('DELETE FROM lodges WHERE id=?', (new,))


def test_restore_coerces_sqlite_values(app_module):
    assert app_module._bk_in('', 'integer') is None and app_module._bk_in(' 12 ', 'bigint') == 12
    assert app_module._bk_in('12.0', 'integer') == 12 and app_module._bk_in(5, 'text') == '5'
    assert app_module._bk_in({'$b64': 'AAE='}, 'bytea') == b'\x00\x01'
