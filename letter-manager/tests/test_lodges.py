# Ενότητα Συμβολικές Στοές: βάση Στοών, καταχώριση, επεξεργασία, εισαγωγή/εξαγωγή.
import re


def test_seed_lodges_listed(admin):
    t = admin.get('/lodges').text
    assert 'LA PAIX' in t and 'ΦΟΙΝΙΞ ΚΕΡΚΥΡΑΣ' in t


def test_create_edit_lodge(admin):
    r = admin.post('/lodges/new', data={'number': '999', 'name': 'ΔΟΚΙΜΗ', 'orient': 'Αθηνών',
                                        'provincial': 'ΕπΜΣτ. Αθηνών', 'email': 'lodge999@example.com',
                                        'status': 'Ενεργή'})
    assert r.status_code in (302, 303), r.text[:300]
    t = admin.get('/lodges').text
    lid = int(re.search(r'/lodges/edit/(\d+)[^<]*</a>(?:(?!</tr>).)*', t, re.S).group(1)) if 'ΔΟΚΙΜΗ' in t else None
    assert lid
    m = re.findall(r'<tr[^>]*>(?:(?!</tr>).)*999(?:(?!</tr>).)*?/lodges/edit/(\d+)', t, re.S)
    lid = int(m[0]) if m else lid
    assert admin.get(f'/lodges/edit/{lid}').status_code == 200


def test_lodge_in_letter_recipients(admin):
    admin.post('/lodges/new', data={'number': '998', 'name': 'ΠΑΡΑΛΗΠΤΡΙΑ', 'email': 'lodge998@example.com',
                                    'provincial': 'ΕπΜΣτ. Αθηνών', 'status': 'Ενεργή'})
    assert 'lodge998@example.com' in admin.get('/new').text


def test_export_import(admin):
    r = admin.get('/lodges/export.xlsx')
    assert r.status_code == 200 and r.content[:2] == b'PK'
    r = admin.post('/lodges/import', data={'pasted': 'Αριθμός\tΌνομα\tEmail\n95\tLA PAIX\tlapaix@example.com'})
    assert r.status_code in (200, 302, 303)
    assert 'lapaix@example.com' in admin.get('/lodges').text


def test_ritual_and_meeting_place(admin, app_module):
    r = admin.post('/lodges/new', data={'number': '996', 'name': 'ΤΥΠΙΚΟΥ', 'ritual': 'Emulation',
                                        'meeting_place': 'Τεκτονικόν Μέγαρον, Ερεσού 38', 'status': 'Ενεργή'})
    assert r.status_code in (302, 303)
    x = [l for l in app_module._lodges_all() if l['number'] == '996'][0]
    assert x['ritual'] == 'Emulation' and x['meeting_place'].startswith('Τεκτονικόν')
    assert 'Emulation' in admin.get('/lodges', params={'q': 'Emulation'}).text
    assert '<option value="Emulation">' in admin.get('/lodges/new').text  # προτάσεις Τυπικού
    r = admin.post('/lodges/import', data={'pasted': 'Αριθμός\tΌνομα\tΤυπικό\tΤόπος συνεδριάσεων\n996\tΤΥΠΙΚΟΥ\tΣκωτικό\tΠειραιάς'})
    assert r.status_code in (302, 303)
    x = [l for l in app_module._lodges_all() if l['number'] == '996'][0]
    assert x['ritual'] == 'Σκωτικό' and x['meeting_place'] == 'Πειραιάς'


def test_export_has_new_columns(admin):
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(admin.get('/lodges/export.xlsx').content)).active
    head = [c.value for c in ws[1]]
    assert 'Τυπικό' in head and 'Τόπος συνεδριάσεων' in head
