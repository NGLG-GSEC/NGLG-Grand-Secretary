# Μητρώο Μελών — πίνακες, αρχική φόρτωση, καταχώριση/αποθήκευση μέλους, Excel ασφαλείας.
# Φορτώνεται από το app.py (βλ. MODULES) στον κοινό χώρο ονομάτων της εφαρμογής.

from io import BytesIO,StringIO

import gzip,csv,unicodedata

from fastapi import UploadFile, File

from openpyxl import Workbook, load_workbook

from openpyxl.styles import Font, PatternFill, Alignment

from openpyxl.utils import get_column_letter

MEMBER_BACKUP_EMAIL='dskiad@gmail.com'

MEMBER_BACKUP_SUBJECT='ΑΠΟΡΡΗΤΟ ΑΡΧΕΙΟ ΕΠΑΦΩΝ ΕΜΤΣΕ ΠΡΟΣ ΔΙΑΓΡΑΦΗ'

MEMBER_SEED_VERSION='2026-09-18-source-1'

def _member_init():
    with con() as c:
        c.executescript("""CREATE TABLE IF NOT EXISTS member_registry(
        id INTEGER PRIMARY KEY AUTOINCREMENT,source_row INTEGER,registry_no INTEGER,surname TEXT NOT NULL DEFAULT '',first_name TEXT NOT NULL DEFAULT '',
        email TEXT DEFAULT '',other_emails TEXT DEFAULT '',mobile TEXT DEFAULT '',other_mobiles TEXT DEFAULT '',degree TEXT DEFAULT '',
        declared_lodge_count TEXT DEFAULT '',deregistered_note TEXT DEFAULT '',additional_lodges TEXT DEFAULT '',active INTEGER NOT NULL DEFAULT 1,
        created_at TEXT,updated_at TEXT);
        CREATE INDEX IF NOT EXISTS ix_member_registry_name ON member_registry(surname,first_name);
        CREATE INDEX IF NOT EXISTS ix_member_registry_email ON member_registry(email);
        CREATE INDEX IF NOT EXISTS ix_member_registry_mobile ON member_registry(mobile);
        CREATE TABLE IF NOT EXISTS member_lodges(
        id INTEGER PRIMARY KEY AUTOINCREMENT,member_id BIGINT NOT NULL,seq INTEGER NOT NULL DEFAULT 1,lodge_name TEXT DEFAULT '',
        lodge_number TEXT DEFAULT '',member_status TEXT DEFAULT '',created_at TEXT,updated_at TEXT,
        FOREIGN KEY(member_id) REFERENCES member_registry(id) ON DELETE CASCADE);
        CREATE INDEX IF NOT EXISTS ix_member_lodges_member ON member_lodges(member_id);
        CREATE INDEX IF NOT EXISTS ix_member_lodges_name ON member_lodges(lodge_name);
        CREATE TABLE IF NOT EXISTS member_degrees_offices(
        id INTEGER PRIMARY KEY AUTOINCREMENT,member_id BIGINT NOT NULL,record_type TEXT DEFAULT '',degree TEXT DEFAULT '',office TEXT DEFAULT '',
        decree_id BIGINT,decree_no INTEGER,decree_year INTEGER,valid_from TEXT DEFAULT '',valid_to TEXT DEFAULT '',is_current INTEGER NOT NULL DEFAULT 1,
        notes TEXT DEFAULT '',created_at TEXT,updated_at TEXT,FOREIGN KEY(member_id) REFERENCES member_registry(id) ON DELETE CASCADE);
        CREATE INDEX IF NOT EXISTS ix_member_offices_member ON member_degrees_offices(member_id);
        CREATE INDEX IF NOT EXISTS ix_member_offices_decree ON member_degrees_offices(decree_id);""")
        if USE_PG:
            c.execute("ALTER TABLE member_registry ADD COLUMN IF NOT EXISTS registry_no INTEGER")
        else:
            mcols=[r['name'] for r in c.execute("PRAGMA table_info(member_registry)")]
            if 'registry_no' not in mcols:c.execute("ALTER TABLE member_registry ADD COLUMN registry_no INTEGER")
        for col in ('surname_variants','first_name_variants'):
            if USE_PG:c.execute(f"ALTER TABLE member_registry ADD COLUMN IF NOT EXISTS {col} TEXT DEFAULT ''")
            elif col not in [r['name'] for r in c.execute("PRAGMA table_info(member_registry)")]:c.execute(f"ALTER TABLE member_registry ADD COLUMN {col} TEXT DEFAULT ''")
        try:c.execute("CREATE UNIQUE INDEX IF NOT EXISTS ix_member_registry_regno ON member_registry(registry_no) WHERE registry_no IS NOT NULL")
        except Exception:pass
    _seed_member_registry()

def _seed_payload():
    env_parts=[]
    for i in range(100):
        v=os.getenv(f'MEMBER_REGISTRY_SEED_{i:03d}','')
        if not v:break
        env_parts.append(v)
    if env_parts:
        try:return json.loads(gzip.decompress(base64.b64decode(''.join(env_parts))).decode('utf-8'))
        except Exception:return None
    raw_parts=sorted((BASE/'static').glob('member_registry_seed.json.part.*'))
    if raw_parts:
        try:return json.loads(''.join(p.read_text(encoding='utf-8') for p in raw_parts))
        except Exception:return None
    fs=sorted((BASE/'static').glob('member_registry_seed.json.gz.b64.*'))
    if not fs:return None
    try:
        s=''.join(p.read_text(encoding='ascii') for p in fs)
        return json.loads(gzip.decompress(base64.b64decode(s)).decode('utf-8'))
    except Exception:return None

def _insert_member(c,m,member_id=None):
    ts=now()
    rn=m.get('registry_no');rn=int(rn) if rn not in (None,'') else None
    vals=(m.get('source_row'),rn,m.get('surname','').strip(),m.get('first_name','').strip(),m.get('email','').strip(),m.get('other_emails','').strip(),
          m.get('mobile','').strip(),m.get('other_mobiles','').strip(),m.get('degree','').strip(),str(m.get('declared_lodge_count','')).strip(),
          m.get('deregistered_note','').strip(),m.get('additional_lodges','').strip(),1 if m.get('active',True) else 0,
          (m.get('surname_variants') or '').strip(),(m.get('first_name_variants') or '').strip(),ts,ts)
    if member_id:
        c.execute("""INSERT INTO member_registry(id,source_row,registry_no,surname,first_name,email,other_emails,mobile,other_mobiles,degree,declared_lodge_count,deregistered_note,additional_lodges,active,surname_variants,first_name_variants,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",(int(member_id),)+vals);mid=int(member_id)
    else:
        cur=c.execute("""INSERT INTO member_registry(source_row,registry_no,surname,first_name,email,other_emails,mobile,other_mobiles,degree,declared_lodge_count,deregistered_note,additional_lodges,active,surname_variants,first_name_variants,created_at,updated_at)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",vals);mid=cur.lastrowid
    for j,l in enumerate(m.get('lodges') or [],1):
        c.execute("""INSERT INTO member_lodges(member_id,seq,lodge_name,lodge_number,member_status,created_at,updated_at) VALUES(?,?,?,?,?,?,?)""",
                  (mid,int(l.get('seq') or j),l.get('name','').strip(),l.get('number','').strip(),l.get('status','').strip(),ts,ts))
    return mid

def _seed_member_registry():
    with con() as c:
        existing=c.execute('SELECT COUNT(*) n FROM member_registry').fetchone()['n']
        flag=c.execute("SELECT value FROM settings WHERE key='member_registry_seed_version'").fetchone()
        if flag and existing:return
        data=_seed_payload()
        if not data:
            print('[member-registry] seed payload unavailable')
            return
        if not existing:
            for m in data.get('members',[]):_insert_member(c,m)
        c.execute("REPLACE INTO settings VALUES(?,?)",('member_registry_seed_version',MEMBER_SEED_VERSION))

def _member_by_id(mid):
    with con() as c:
        r=c.execute('SELECT * FROM member_registry WHERE id=?',(mid,)).fetchone()
        if not r:return None,[]
        ls=[dict(x) for x in c.execute('SELECT * FROM member_lodges WHERE member_id=? ORDER BY seq,id',(mid,))]
        return dict(r),ls

def _xlsx_bytes():
    with con() as c:
        ms=[dict(r) for r in c.execute('SELECT * FROM member_registry ORDER BY surname,first_name,id')]
        ls=[dict(r) for r in c.execute('SELECT * FROM member_lodges ORDER BY member_id,seq,id')]
        os=[dict(r) for r in c.execute('SELECT * FROM member_degrees_offices ORDER BY member_id,id')]
    wb=Workbook();ws=wb.active;ws.title='ΜΗΤΡΩΟ ΜΕΛΩΝ'
    ws.append(['ID','Αρ. Μητρώου','Επώνυμο','Όνομα','Κύριο Email','Άλλα Email','Κύριο Κινητό','Άλλα Κινητά','Τεκτονικός Βαθμός','Ενεργός','Σημείωση Διαγραφής','Δηλωμένος Αρ. Στοών','Πρόσθετες Στοές','Γραμμή Πηγής','Ενημερώθηκε','Παραλλαγές Επωνύμου','Παραλλαγές Ονόματος'])
    for m in ms:ws.append([m['id'],m.get('registry_no'),m['surname'],m['first_name'],m['email'],m['other_emails'],m['mobile'],m['other_mobiles'],m['degree'],'ΝΑΙ' if m['active'] else 'ΟΧΙ',m['deregistered_note'],m['declared_lodge_count'],m['additional_lodges'],m['source_row'],m['updated_at'],m.get('surname_variants') or '',m.get('first_name_variants') or ''])
    sl=wb.create_sheet('ΣΤΟΕΣ ΜΕΛΩΝ');sl.append(['ID Μέλους','Α/Α Στοάς','Στοά','Αριθμός Στοάς','Κατάσταση'])
    for l in ls:sl.append([l['member_id'],l['seq'],l['lodge_name'],l['lodge_number'],l['member_status']])
    so=wb.create_sheet('ΒΑΘΜΟΙ & ΑΞΙΩΜΑΤΑ');so.append(['ID','ID Μέλους','Τύπος Εγγραφής','Βαθμός','Αξίωμα','ID Διατάγματος','Αρ. Διατάγματος','Έτος','Από','Έως','Ενεργό','Σημειώσεις'])
    for o in os:so.append([o['id'],o['member_id'],o['record_type'],o['degree'],o['office'],o['decree_id'],o['decree_no'],o['decree_year'],o['valid_from'],o['valid_to'],'ΝΑΙ' if o['is_current'] else 'ΟΧΙ',o['notes']])
    for sh in (ws,sl,so):
        sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1F4E78');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        widths={}
        for row in sh.iter_rows():
            for cell in row:
                cell.alignment=Alignment(vertical='top',wrap_text=True);v='' if cell.value is None else str(cell.value);widths[cell.column]=min(max(widths.get(cell.column,0),len(v)+2),42)
        for col,w in widths.items():sh.column_dimensions[get_column_letter(col)].width=max(10,w)
    b=BytesIO();wb.save(b);return b.getvalue()

def _send_member_backup(data,reason='διαγραφή'):
    host=os.getenv('SMTP_HOST');usr=os.getenv('SMTP_USERNAME');pwd=os.getenv('SMTP_PASSWORD');port=int(os.getenv('SMTP_PORT','587'))
    if not(host and usr and pwd):raise HTTPException(503,'Η διαγραφή δεν εκτελέστηκε: δεν είναι διαθέσιμη η υποχρεωτική αποστολή Excel ασφαλείας.')
    m=EmailMessage();m['Subject']=MEMBER_BACKUP_SUBJECT;m['From']=os.getenv('OTP_FROM_EMAIL',usr);m['To']=MEMBER_BACKUP_EMAIL
    m.set_content('Αυτόματο πλήρες αντίγραφο του Μητρώου Μελών πριν από '+reason+'.\nΗ ενέργεια δημιουργήθηκε από την Ψηφιακή Μεγάλη Γραμματεία.')
    fn='EMSTE_MEMBER_REGISTRY_PRE_DELETE_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.xlsx'
    m.add_attachment(data,maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename=fn)
    try:
        with smtplib.SMTP(host,port,timeout=30) as s:s.starttls(context=ssl.create_default_context());s.login(usr,pwd);s.send_message(m)
    except Exception:raise HTTPException(503,'Η διαγραφή δεν εκτελέστηκε επειδή απέτυχε η υποχρεωτική αποστολή του Excel ασφαλείας.')

def _lodge_text(ls):return '\n'.join(' | '.join([x.get('lodge_name',''),x.get('lodge_number',''),x.get('member_status','')]).rstrip(' |') for x in ls)

def _parse_lodges(v):
    out=[]
    for i,line in enumerate((v or '').splitlines(),1):
        if not line.strip():continue
        p=[x.strip() for x in line.split('|')]
        while len(p)<3:p.append('')
        out.append({'seq':i,'name':p[0],'number':p[1],'status':' | '.join(p[2:]).strip()})
    return out

def _member_form(m=None,ls=None):
    m=m or {};ls=ls or []
    return f"""<div class="grid card">
<div><label>Αρ. Μητρώου</label><input name="registry_no" value="{esc(str(m.get('registry_no') or ''))}"></div>
<div><label>Επώνυμο</label><input name="surname" value="{esc(m.get('surname',''))}" required></div>
<div><label>Όνομα</label><input name="first_name" value="{esc(m.get('first_name',''))}" required></div>
<div><label>Παραλλαγές Επωνύμου</label><input name="surname_variants" value="{esc(m.get('surname_variants') or '')}" placeholder="π.χ. CASTANEDA; ΚΑΣΤΑΝΕΔΑ"></div>
<div><label>Παραλλαγές Ονόματος</label><input name="first_name_variants" value="{esc(m.get('first_name_variants') or '')}" placeholder="π.χ. CARLOS; ΚΑΡΛΟΣ"></div>
<div><label>Κύριο Email</label><input type="email" name="email" value="{esc(m.get('email',''))}"></div>
<div><label>Άλλα Email</label><input name="other_emails" value="{esc(m.get('other_emails',''))}"></div>
<div><label>Κύριο Κινητό</label><input name="mobile" value="{esc(m.get('mobile',''))}"></div>
<div><label>Άλλα Κινητά</label><input name="other_mobiles" value="{esc(m.get('other_mobiles',''))}"></div>
<div><label>Τεκτονικός Βαθμός</label><input name="degree" value="{esc(m.get('degree',''))}"></div>
<div><label>Ενεργός</label><select name="active"><option value="1" {'selected' if m.get('active',1) else ''}>ΝΑΙ</option><option value="0" {'selected' if not m.get('active',1) else ''}>ΟΧΙ</option></select></div>
<div class="full"><label>Σημείωση Διαγραφής</label><input name="deregistered_note" value="{esc(m.get('deregistered_note',''))}"></div>
<div class="full"><label>Πρόσθετες Στοές / παλαιά πληροφορία</label><textarea name="additional_lodges">{esc(m.get('additional_lodges',''))}</textarea></div>
<div class="full"><label>Στοές</label><textarea name="lodges_text" placeholder="Μία στοά ανά γραμμή: ΟΝΟΜΑ | ΑΡΙΘΜΟΣ | ΚΑΤΑΣΤΑΣΗ">{esc(_lodge_text(ls))}</textarea><small>Παράδειγμα: ΠΑΡΘΕΝΩΝ | 3 | 1. ΤΑΚΤΙΚΟ</small></div>
</div>"""

def _save_member(c,f,mid=None):
    m={k:str(f.get(k,'')).strip() for k in ['surname','first_name','email','other_emails','mobile','other_mobiles','degree','deregistered_note','additional_lodges','surname_variants','first_name_variants']}
    m['active']=str(f.get('active','1'))=='1';m['lodges']=_parse_lodges(str(f.get('lodges_text','')));m['declared_lodge_count']=str(len(m['lodges']))
    rn=str(f.get('registry_no','')).strip();m['registry_no']=int(rn) if rn else None
    if not m['surname'] or not m['first_name']:raise HTTPException(400,'Συμπληρώστε επώνυμο και όνομα.')
    if not mid:return _insert_member(c,m)
    ts=now();c.execute("""UPDATE member_registry SET registry_no=?,surname=?,first_name=?,email=?,other_emails=?,mobile=?,other_mobiles=?,degree=?,declared_lodge_count=?,deregistered_note=?,additional_lodges=?,active=?,surname_variants=?,first_name_variants=?,updated_at=? WHERE id=?""",
        (m['registry_no'],m['surname'],m['first_name'],m['email'],m['other_emails'],m['mobile'],m['other_mobiles'],m['degree'],m['declared_lodge_count'],m['deregistered_note'],m['additional_lodges'],1 if m['active'] else 0,m['surname_variants'],m['first_name_variants'],ts,mid))
    c.execute('DELETE FROM member_lodges WHERE member_id=?',(mid,))
    for l in m['lodges']:c.execute("""INSERT INTO member_lodges(member_id,seq,lodge_name,lodge_number,member_status,created_at,updated_at) VALUES(?,?,?,?,?,?,?)""",(mid,l['seq'],l['name'],l['number'],l['status'],ts,ts))
    return mid
